import docx
from docx.shared import Pt
import os

# Create document
doc = docx.Document()

# Add Title
title = doc.add_heading('Step-by-Step Instructions: Expanded Wine Market Analysis', 0)

# 1. Environment Setup
doc.add_heading('1. Environment Setup', level=1)
doc.add_paragraph('Ensure you have Python installed (3.8+ recommended). Install the required libraries using the following command in your terminal:')
doc.add_paragraph('pip install pandas matplotlib seaborn python-docx nbformat wordcloud')

# 2. Project Files
doc.add_heading('2. Project Phase Scripts', level=1)
doc.add_paragraph('The following scripts represent the stages of the investigation:')
files = [
    'clean_data.py: Preprocessing & Dataset Merger',
    'analysis_correlation.py: Price vs. Quality Analysis',
    'analysis_market.py: Market Share & Geography',
    'analysis_tasters.py: Critic Score Consistency',
    'analysis_price_segments.py: Price Tiering (Budget to Luxury)',
    'analysis_vintage_extraction.py: Year Data Extraction',
    'analysis_heatmap_geo.py: Regional Specialization Heatmap',
    'analysis_wordcloud_high.py: High-Rating Descriptor Profile',
    'analysis_wordcloud_low.py: Low-Rating Descriptor Profile',
    'analysis_top_value_bottles.py: Points-per-Dollar ROI Study',
    'analysis_winery_stars.py: Identifying Leading Producers',
    'generate_final_expanded_notebook.py: Final Report Consolidation'
]
for f in files:
    doc.add_paragraph(f, style='List Bullet')

# 3. Step-by-Step Investigation Flow
doc.add_heading('3. Investigation Flow (12 Terminal Interactions)', level=1)
doc.add_paragraph('Run these scripts in sequence to replicate the full Investigation:')

steps = [
    'python clean_data.py (Initial Data Prep)',
    'python analysis_correlation.py (Initial Correlations)',
    'python analysis_market.py (Variety & Country Volume)',
    'python analysis_tasters.py (Taster Bias Check)',
    'python analysis_price_segments.py (Market Tiering)',
    'python analysis_vintage_extraction.py (Vintage Analysis)',
    'python analysis_heatmap_geo.py (Geo-Variety Mapping)',
    'python analysis_wordcloud_high.py (Linguistic profiling - High)',
    'python analysis_wordcloud_low.py (Linguistic profiling - Low)',
    'python analysis_top_value_bottles.py (Value Opportunity List)',
    'python analysis_winery_stars.py (Star Producer Identification)',
    'python generate_final_expanded_notebook.py (Creates MiraclePython.ipynb)'
]

for i, step in enumerate(steps, 1):
    doc.add_paragraph(f'{i}. {step}')

# 4. Final Deliverables
doc.add_heading('4. Final Deliverables', level=1)
doc.add_paragraph('Upon completion, you will find the comprehensive analysis and assortment strategy in:')
doc.add_paragraph('MiraclePython.ipynb (Open with Jupyter or VS Code)')

# Save document
output_path = r'c:\Users\E-BOOKS WORKSTATION\Desktop\miracle\How_To_Run_Analysis.docx'
doc.save(output_path)

print(f"Documentation updated and saved to {output_path}")
